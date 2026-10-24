import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def create_report():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header / University
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni1 = p_uni.add_run("សាកលវិទ្យាល័យ សៅស៍អុីសថ៍អេយេសៀ\n")
    r_uni1.font.name = 'Khmer OS Muol Light'
    r_uni1.font.size = Pt(13)
    r_uni1.bold = True
    r_uni1.font.color.rgb = RGBColor(16, 44, 87)

    r_uni2 = p_uni.add_run("UNIVERSITY OF SOUTH-EAST ASIA\n")
    r_uni2.font.name = 'Arial'
    r_uni2.font.size = Pt(12)
    r_uni2.bold = True
    r_uni2.font.color.rgb = RGBColor(16, 44, 87)

    r_fac = p_uni.add_run("Faculty of Science and Technology\n\n")
    r_fac.font.name = 'Arial'
    r_fac.font.size = Pt(11)
    r_fac.font.color.rgb = RGBColor(71, 85, 105)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_hw = p_title.add_run("Homework 3\n")
    r_hw.font.name = 'Arial'
    r_hw.font.size = Pt(18)
    r_hw.bold = True
    r_hw.font.color.rgb = RGBColor(37, 99, 235)

    r_t = p_title.add_run("RELEASE MANAGEMENT: GITHUB TAGS & JENKINS BUILD/PUSH/DEPLOY/ROLLBACK\n")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(13)
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(30, 41, 59)

    # Info Box Table
    info_table = doc.add_table(rows=6, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Course:", "Advanced Cloud Technology"),
        ("Lecturer:", "Mr. Krol Noeurn"),
        ("Student Name:", "TONN Sievphor"),
        ("GitHub Repository:", "https://github.com/Sievphor/cloud_home_work_3.git"),
        ("Amazon ECR URI:", "908157891706.dkr.ecr.us-east-1.amazonaws.com/sievphor-homework3"),
        ("Academic Year:", "2026 - 2027")
    ]
    for i, (k, v) in enumerate(info_data):
        row = info_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.text = k
        c1.text = v
        c0.paragraphs[0].runs[0].font.bold = True
        c0.paragraphs[0].runs[0].font.name = 'Arial'
        c0.paragraphs[0].runs[0].font.size = Pt(10)
        c1.paragraphs[0].runs[0].font.name = 'Arial'
        c1.paragraphs[0].runs[0].font.size = Pt(10)
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")

    doc.add_page_break()

    def add_section_h1(text):
        h = doc.add_paragraph()
        r = h.add_run(text)
        r.font.name = 'Khmer OS Battambang'
        r.font.size = Pt(12)
        r.bold = True
        r.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_body_p(text):
        p = doc.add_paragraph()
        r = p.add_run(text)
        r.font.name = 'Khmer OS Battambang'
        r.font.size = Pt(10)
        p.paragraph_format.line_spacing = 1.15
        return p

    def add_figure_box(title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"[{title} - សូម Paste រូបភាព Screenshot នៅទីនេះ]")
        r.font.name = 'Khmer OS Battambang'
        r.font.size = Pt(9.5)
        r.italic = True
        r.font.color.rgb = RGBColor(100, 116, 139)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(12)

    # 1. Introduction
    add_section_h1("១. សេចក្តីផ្តើម (Introduction)")
    add_body_p(
        "កិច្ចការ Homework 3 ផ្តោតលើការគ្រប់គ្រងការបញ្ចេញកូដ (Release Management) ដោយប្រើប្រាស់ GitHub Tags ជាមួយ "
        "Semantic Versioning (vMAJOR.MINOR.PATCH)។ ជំនួសឱ្យការ Deploy រាល់ពេលមាន Commit តូចតាច ប្រព័ន្ធ CI/CD របស់ Jenkins "
        "ដំណើរការតែនៅពេលមានការ Push Git Tag (v1.0.0, v1.1.0, v1.2.0) ប៉ុណ្ណោះ។ Docker Image ទាំងអស់ត្រូវបានដាក់ Tag "
        "ដូចគ្នានឹង Git Tag និងរក្សាទុកជា Immutable Tags ក្នុង Amazon ECR។ លើសពីនេះ ប្រព័ន្ធត្រូវបានបំពាក់នូវយន្តការ Rollback "
        "ដែលអនុញ្ញាតឱ្យត្រឡប់ទៅកាន់កំណែចាស់ពីមុនបានភ្លាមៗ ដោយមិនបាច់ Build Image ឡើងវិញឡើយ (No Rebuild)។"
    )

    # 2. Objectives & Repository
    add_section_h1("២. គោលបំណង និងរចនាសម្ព័ន្ធគម្រោង (Objectives & Repository)")
    add_body_p(
        "• អនុវត្ត Semantic Versioning (v1.0.0, v1.1.0, v1.2.0) ក្នុងការគ្រប់គ្រង Release\n"
        "• កំណត់ Jenkins Pipeline ឱ្យរត់តាម GitHub Tag Webhook (Tag-only trigger)\n"
        "• ធានាថា Docker Image Tag ត្រូវគ្នានឹង Git Tag និងរក្សាទុកគ្រប់រូបភាពទាំងអស់ក្នុង ECR\n"
        "• Parameterize Pipeline ជាមួយ ACTION (all, build, push, deploy, rollback) និង VERSION\n"
        "• អនុវត្ត Rollback ដោយទាញយក Image ចាស់ពី ECR មក Deploy លើ Docker Swarm ដោយគ្មាន Rebuild។\n"
        "គម្រោងត្រូវបានរៀបចំឡើងយ៉ាងស្អាតនៅលើ GitHub: https://github.com/Sievphor/cloud_home_work_3 រួមមានឯកសារ RELEASING.md, Jenkinsfile, docker-stack.yml និងកូដគម្រោង។"
    )
    add_figure_box("រូបភាពទី ១៖ GitHub Repository & RELEASING.md (cloud_home_work_3)")

    # 3. Implementation Steps
    add_section_h1("៣. ដំណើរការអនុវត្តជាក់ស្តែង (Implementation Steps)")
    add_body_p(
        "៣.១ ការបង្កើត Git Tags តាម Semantic Versioning:\n"
        "បានបង្កើត Releases ចំនួន ៣ រួមមាន v1.0.0, v1.1.0, និង v1.2.0 នៅក្នុង GitHub Repository "
        "https://github.com/Sievphor/cloud_home_work_3.git ដោយប្រើពាក្យបញ្ជា 'git tag -a <tag> -m <msg>' "
        "និងរុញទៅ GitHub តាមរយៈ 'git push origin --tags'។"
    )
    add_figure_box("រូបភាពទី ២៖ GitHub Releases & Tags (v1.0.0, v1.1.0, v1.2.0)")

    add_body_p(
        "៣.២ ការកំណត់ Jenkins Parameterized Pipeline:\n"
        "Jenkinsfile ត្រូវបានរៀបចំឱ្យទទួល Parameters ACTION (all, build, push, deploy, rollback) និង VERSION។ "
        "នៅពេលមាន Tag Push មកដល់ Jenkins នឹងដំណើរការ Build និង Deploy ដោយស្វ័យប្រវត្តិ។"
    )
    add_figure_box("រូបភាពទី ៣៖ Jenkins Parameterized Build (ACTION & VERSION)")

    add_body_p(
        "៣.៣ ការរក្សាទុក Image ជា Immutable Tags ក្នុង Amazon ECR:\n"
        "រូបភាពនីមួយៗត្រូវបាន Push ទៅកាន់ ECR Repository 'sievphor-homework3' "
        "(URI: 908157891706.dkr.ecr.us-east-1.amazonaws.com/sievphor-homework3)។ "
        "រូបភាព v1.0.0, v1.1.0, និង v1.2.0 ត្រូវបានរក្សាទុកទាំងអស់ជា Immutable Tags ដោយមិនមានការសរសេរជាន់លើឡើយ។"
    )
    add_figure_box("រូបភាពទី ៤៖ Amazon ECR Immutable Tags (v1.0.0, v1.1.0, v1.2.0)")

    add_body_p(
        "៣.៤ ការ Deploy លើ Docker Swarm និងការបង្ហាញវេបសាយ:\n"
        "Jenkins ធ្វើបច្ចុប្បន្នភាពលើ Service homework3_web លើ Port 80 តាមរយៈ Rolling Update "
        "ធានាថាកម្មវិធីនៅតែដំណើរការធម្មតាគ្រប់ពេលវេលាលើ Cluster Nodes ទាំង ៣ (Manager, Worker 1, Worker 2)។"
    )
    add_figure_box("រូបភាពទី ៥៖ វេបសាយ Version v1.0.0 (Port 80) បង្ហាញ Badge Version")

    # 4. Rollback
    add_section_h1("៤. ការអនុវត្ត Rollback ដោយមិនបាច់ Rebuild (Zero Rebuild Rollback)")
    add_body_p(
        "នៅពេលរកឃើញបញ្ហាក្នុងកំណែ v1.2.0 យើងបានធ្វើការ Rollback ត្រឡប់ទៅកំណែចាស់ v1.0.0៖\n"
        "១. ចូលទៅកាន់ Jenkins Job ហើយជ្រើសរើស ACTION = rollback និង VERSION = v1.0.0\n"
        "២. Pipeline បានផ្ទៀងផ្ទាត់វត្តមានរបស់ Image v1.0.0 នៅក្នុង Amazon ECR ជាមុនតាមរយៈ aws ecr describe-images\n"
        "៣. Jenkins បានបញ្ជាទៅ Swarm Manager (3.91.30.20) ឱ្យធ្វើ Rolling Update ទៅកាន់ Image v1.0.0 ដោយផ្ទាល់\n"
        "៤. ដំណើរការនេះពុំមានការ Checkout កូដ ឬ Build រូបភាពថ្មីឡើយ (No Rebuild) ដែលជួយកាត់បន្ថយពេលវេលា និងគ្មានហានិភ័យនៃកំហុសកូដ។"
    )
    add_figure_box("រូបភាពទី ៦៖ ការអនុវត្ត Rollback ក្នុង Jenkins (No Rebuild Console Log)")
    add_figure_box("រូបភាពទី ៧៖ ស្ថានភាព Swarm Service ក្រោយ Rollback (docker service inspect)")
    add_figure_box("រូបភាពទី ៨៖ វេបសាយក្រោយ Rollback ជោគជ័យ (Port 80 ដំណើរការធម្មតា)")

    # 5. History Table
    add_section_h1("៥. កំណត់ត្រាប្រវត្តិនៃការបញ្ចេញកូដ និង Rollback (Deployment History)")
    table = doc.add_table(rows=5, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Version", "Action", "Image Source", "Result"]
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.name = 'Arial'
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)
        set_cell_background(cell, "E2E8F0")

    history_data = [
        ("v1.0.0", "Deploy Release", "Built & Pushed to ECR", "Success (Replicas 3/3)"),
        ("v1.1.0", "Deploy Release", "Built & Pushed to ECR", "Success (Rolling Update)"),
        ("v1.2.0", "Deploy Release", "Built & Pushed to ECR", "Success (Rolling Update)"),
        ("v1.0.0", "Rollback", "Direct from ECR (No Rebuild)", "Rollback Verified Success")
    ]
    for i, row_data in enumerate(history_data):
        row = table.rows[i + 1]
        for j, val in enumerate(row_data):
            cell = row.cells[j]
            cell.text = val
            cell.paragraphs[0].runs[0].font.name = 'Arial'
            cell.paragraphs[0].runs[0].font.size = Pt(9)
            set_cell_background(cell, "F8FAFC" if i % 2 == 0 else "FFFFFF")

    # 6. Conclusion
    add_section_h1("៦. សេចក្តីសន្និដ្ឋាន (Conclusion)")
    add_body_p(
        "កិច្ចការ Homework 3 នេះបានបង្ហាញពីប្រសិទ្ធភាពខ្ពស់នៃ Release Management ក្នុងប្រព័ន្ធ Cloud DevOps។ "
        "ការរួមបញ្ចូលគ្នារវាង Semantic Versioning, Immutable Image Tags ក្នុង Amazon ECR, Parameterized Jenkins Pipeline, "
        "និង Zero-Downtime Rollback លើ Docker Swarm ផ្តល់នូវស្តង់ដារវិស្វកម្មកម្រិត Production ដ៏រឹងមាំសម្រាប់គ្រប់គ្រងកម្មវិធី។\n"
        "• Manager Endpoint: http://3.91.30.20:80\n"
        "• Worker 1 Endpoint: http://34.229.205.236:80\n"
        "• Worker 2 Endpoint: http://54.160.200.116:80"
    )

    output_path = os.path.join(os.path.dirname(__file__), "TONN-Sievphor-Documentation_homework3.docx")
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    create_report()
